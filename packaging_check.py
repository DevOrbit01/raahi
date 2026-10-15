"""Offline checks of the distributed executable; no browser or live scrape."""
from datetime import datetime, timezone
import json
from pathlib import Path
import platform
import sys
from tempfile import TemporaryDirectory
import traceback


def run_packaging_check(report_path):
    report = {
        'passed': False,
        'platform': sys.platform,
        'architecture': platform.machine(),
        'python': platform.python_version(),
        'frozen': bool(getattr(sys, 'frozen', False)),
        'tested_at_utc': datetime.now(timezone.utc).isoformat(),
        'checks': {},
        'live_cloudflare_tested': False,
    }
    app = None
    try:
        from app_paths import application_data_dir, default_output_folder
        from gui import ScraperGUI
        from openpyxl import load_workbook
        from playwright.sync_api import sync_playwright
        from cloakbrowser import binary_info
        from scraper_runner import save_items_to_excel
        from scrapers.local_browser import _profile_lock, browser_backend, browser_ready
        report['checks']['per_user_paths'] = {
            'app_data': application_data_dir(), 'exports': default_output_folder(),
        }
        with sync_playwright() as driver:
            assert driver.chromium.name == 'chromium'
            report['checks']['bundled_playwright_driver'] = True
        assert browser_backend() == 'cloak'
        assert browser_ready()
        report['checks']['cloakbrowser_wrapper'] = True
        report['checks']['cloakbrowser_binary_installed_for_runner'] = bool(binary_info()['installed'])
        # The proprietary browser binary must be downloaded by the end user
        # from CloakHQ; it is never redistributed inside Raahi's executable.
        bundle_root = Path(getattr(sys, '_MEIPASS', Path(__file__).parent))
        assert not any(bundle_root.rglob('chrome.exe'))
        report['checks']['cloakbrowser_binary_not_bundled'] = True
        app = ScraperGUI()
        app.root.withdraw()
        app.root.update_idletasks()
        assert len(app.site_vars) == 7 and len(app.state_vars) == 36
        assert app.output_folder == default_output_folder()
        assert app.eauctions_access_btn.winfo_exists()
        assert app.eauction_backend == 'cloak'
        report['checks']['gui_initialized'] = True
        assert not hasattr(app, 'open_browser_connector_setup')
        report['checks']['no_browser_extension_required'] = True
        with TemporaryDirectory() as folder:
            with _profile_lock(Path(folder)/'test-profile'):
                report['checks']['platform_file_lock'] = True
            file = save_items_to_excel([
                {'listingId': 'build-check', 'name': 'Build validation',
                 'images': 'https://www.eauctionsindia.com/storage/images/example.webp'}
            ], output_folder=folder)
            workbook = load_workbook(file)
            assert workbook.active.max_row == 2
            assert 'listingId' in [cell.value for cell in workbook.active[1]]
            workbook.close()
            report['checks']['excel_roundtrip'] = True

            # Exercise the browser scraper from inside the frozen executable,
            # with deterministic pages and no live website dependency.
            from contextlib import nullcontext
            from types import SimpleNamespace
            from unittest.mock import patch
            from scrapers import eauctionsindia
            import gui
            first = eauctionsindia._listing_urls('Delhi', 1)[0]
            second = eauctionsindia._listing_urls('Delhi', 2)[0]
            missing = eauctionsindia.BASE_URL + '/properties/1'
            valid = eauctionsindia.BASE_URL + '/properties/2'
            pages = {
                first: (f'<a href="{missing}">Missing date</a>'
                        f'<ul class="pagination"><a href="{second}" rel="next">Next</a></ul>'),
                second: f'<a href="{valid}">Valid auction</a>',
                missing: '<h1>Incomplete auction</h1><p>Auction ID: 1</p>',
                valid: ('<h1>Complete auction</h1><p>Auction ID: 2</p>'
                        '<div><strong>Province/State:</strong>Delhi</div>'
                        '<div><strong>Auction Start Date:</strong>20-11-2026 12:00 PM</div>'),
            }
            with patch.dict('os.environ', {
                'EAUCTIONSINDIA_BROWSER_BACKEND': 'cloak',
                'EAUCTIONSINDIA_BROWSER_PROFILE': str(Path(folder) / 'skip-test'),
            }), patch.object(eauctionsindia, 'desktop_browser', return_value=nullcontext(
                SimpleNamespace(wait_for_timeout=lambda delay: None)
            )), patch.object(eauctionsindia, '_read_page', side_effect=lambda page, url, *a, **kw: pages[url]):
                items = eauctionsindia.scrape('Delhi')
            assert [item['listingId'] for item in items] == ['2']
            assert len(items.warnings) == 1 and missing in items.warnings[0]
            report['checks']['missing_date_continues_pagination'] = True

            def checked_items(*args, **kwargs):
                kwargs['run_details'].update({
                    'failed_tasks': 0, 'skipped_tasks': 0, 'skipped_properties': 1,
                    'errors': [], 'warnings': items.warnings,
                })
                return items

            with patch.object(gui, 'collect_new_items', side_effect=checked_items):
                app.run_scraping_thread(['Delhi'], ['eauctionsindia'])
            app.root.update()
            assert '1 incomplete listings skipped' in app.status_label.cget('text')
            assert missing in app.log_text.get('1.0', 'end')
            report['checks']['gui_displays_skipped_listing_warning'] = True
        report['passed'] = True
    except Exception:
        report['error'] = traceback.format_exc()
    finally:
        if app is not None:
            app.root.destroy()
        file = Path(report_path)
        file.parent.mkdir(parents=True, exist_ok=True)
        file.write_text(json.dumps(report, indent=2), encoding='utf-8')
    return report['passed']
