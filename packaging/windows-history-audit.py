"""Test the existing release without rebuilding or changing its application code."""
import argparse
import copy
import hashlib
import importlib.abc
import importlib.util
import json
import marshal
import os
from pathlib import Path
import subprocess
import sys
from unittest.mock import patch

EXPECTED_SHA = '48234cba84608196377f9ab825599e921e80d86e6f176bf091f93b2647d499c0'


class ReleaseImporter(importlib.abc.MetaPathFinder, importlib.abc.Loader):
    """Load application modules from the shipped EXE, not the checkout."""
    def __init__(self, exe):
        from PyInstaller.archive.readers import CArchiveReader
        self.archive = CArchiveReader(str(exe))
        self.pyz = self.archive.open_embedded_archive('PYZ.pyz')

    def find_spec(self, fullname, path=None, target=None):
        if fullname == 'gui' or fullname.split('.')[0] in {
            'config', 'fingerprint', 'seen_db', 'scraper_runner', 'app_paths',
            'scrapers', 'packaging_check',
        }:
            return importlib.util.spec_from_loader(fullname, self, is_package=fullname == 'scrapers')
        return None

    def create_module(self, spec):
        return None

    def exec_module(self, module):
        code = (marshal.loads(self.archive.extract('gui')) if module.__name__ == 'gui'
                else self.pyz.extract(module.__name__))
        module.__file__ = 'shipped-exe/' + module.__name__.replace('.', '/') + '.py'
        exec(code, module.__dict__)


def worker(exe, root, phase):
    os.environ['LOCALAPPDATA'] = str(root / 'LocalAppData')
    os.environ['EAUCTIONSINDIA_BROWSER_BACKEND'] = 'cloak'
    sys.meta_path.insert(0, ReleaseImporter(exe))
    import gui
    import scraper_runner
    from seen_db import seen_db

    items = [
        {'source': 'baanknet', 'listingId': f'ASSAM-AUDIT-{i}', 'state': 'Assam',
         'city': 'Test city', 'name': f'Test property {i}', 'date': '15/10/2026',
         'reservePrice': (200000 if phase == 'restart' and i < 2 else 100000)}
        for i in range(3)
    ]
    app = gui.ScraperGUI()
    app.root.withdraw()
    app.deselect_all_sites()
    app.site_vars['baanknet'].set(True)
    app.deselect_all_states()
    app.state_vars['Assam'].set(True)
    app.download_images_var.set(False)
    app.output_folder_entry.delete(0, 'end')
    app.output_folder_entry.insert(0, str(root / 'Exports'))
    result = {'phase': phase, 'platform': sys.platform, 'history_path': seen_db.db_file}
    failures = []

    def fail(exc):
        failures.append(repr(exc))
        app.root.quit()

    def after_scrape():
        if app.is_scraping:
            app.root.after(30, after_scrape)
            return
        try:
            ids = [item['listingId'] for item in app.scraped_items]
            result['listed_ids'] = ids
            if phase == 'restart':
                assert ids == ['ASSAM-AUDIT-2'], ids
                assert len(seen_db.seen) == 2
                result['downloaded_2_skipped_after_process_restart'] = True
                result['unexported_property_remains_available'] = True
                result['changed_price_does_not_reinclude_old_ids'] = True
                app.root.quit()
                return
            assert len(ids) == 3 and len(app.property_check_vars) == 3
            for var, item in app.property_check_vars:
                var.set(item['listingId'] != 'ASSAM-AUDIT-2')
            app.download_selected()
            app.root.after(30, lambda: after_export(1))
        except Exception as exc:
            fail(exc)

    def after_export(number):
        if app.is_downloading:
            app.root.after(30, lambda: after_export(number))
            return
        try:
            from openpyxl import load_workbook
            assert app.output_file and Path(app.output_file).is_file()
            book = load_workbook(app.output_file, read_only=True)
            assert book.active.max_row == 3
            book.close()
            assert len(seen_db.seen) == 2
            assert len(app.property_check_vars) == 3
            if number == 1:
                result['first_export_rows'] = 2
                result['first_export_file'] = app.output_file
                app.download_selected()
                app.root.after(30, lambda: after_export(2))
            else:
                assert result['first_export_file'] != app.output_file
                result['repeat_export_allowed_before_new_scrape'] = True
                result['history_count_after_repeat_export'] = len(seen_db.seen)
                app.root.quit()
        except Exception as exc:
            fail(exc)

    def begin():
        app.start_scraping()
        app.root.after(30, after_scrape)

    with patch.object(scraper_runner.baanknet, 'scrape', side_effect=lambda *a, **kw: copy.deepcopy(items)), \
            patch.object(gui.messagebox, 'showinfo'), \
            patch.object(gui.messagebox, 'showwarning') as warnings, \
            patch.object(gui.messagebox, 'showerror') as errors:
        app.root.after(0, begin)
        app.root.after(60000, lambda: fail(TimeoutError('GUI audit timed out')))
        app.root.mainloop()
        result['warning_count'] = warnings.call_count
        result['error_count'] = errors.call_count
    app.root.destroy()
    result['failures'] = failures
    result['passed'] = not failures and result['warning_count'] == 0 and result['error_count'] == 0
    (root / f'{phase}.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
    print(json.dumps(result, indent=2))
    if not result['passed']:
        raise SystemExit(1)


def audit(exe, root):
    root.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(exe.read_bytes()).hexdigest()
    assert digest == EXPECTED_SHA, digest
    assert sys.platform == 'win32', 'This audit must run on Windows.'
    environment = dict(os.environ, LOCALAPPDATA=str(root / 'NativeSelfTestAppData'))
    native_report = root / 'native-exe-self-test.json'
    native = subprocess.run([str(exe), '--self-test', '--report', str(native_report)],
                            env=environment, timeout=150)
    assert native.returncode == 0
    native_result = json.loads(native_report.read_text(encoding='utf-8'))
    assert native_result['passed'] and native_result['frozen'] and native_result['platform'] == 'win32'
    for phase in ['first', 'restart']:
        if phase == 'restart':
            for path in (root / 'Exports').glob('*.xlsx'):
                path.unlink()
        completed = subprocess.run([
            sys.executable, str(Path(__file__).resolve()), '--exe', str(exe),
            '--root', str(root), '--phase', phase,
        ], timeout=90)
        assert completed.returncode == 0, phase
    first = json.loads((root / 'first.json').read_text(encoding='utf-8'))
    restart = json.loads((root / 'restart.json').read_text(encoding='utf-8'))
    report = {
        'passed': first['passed'] and restart['passed'],
        'execution_platform': sys.platform,
        'executable_sha256': digest,
        'application_rebuilt': False,
        'native_exe_self_test_passed': True,
        'history_tests_use_modules_extracted_from_actual_exe': True,
        'history_test_input': 'Simulated Assam BaankNet properties; no live website request',
        'real_tk_gui_and_download_worker_used': True,
        'first_process': first, 'fresh_process': restart,
        'excel_files_deleted_before_fresh_process': True,
        'client_machine_tested': False,
    }
    (root / 'windows-history-audit.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--exe', required=True, type=lambda value: Path(value).resolve())
    parser.add_argument('--root', required=True, type=lambda value: Path(value).resolve())
    parser.add_argument('--phase', choices=['first', 'restart'])
    arguments = parser.parse_args()
    if arguments.phase:
        worker(arguments.exe, arguments.root, arguments.phase)
    else:
        audit(arguments.exe, arguments.root)
